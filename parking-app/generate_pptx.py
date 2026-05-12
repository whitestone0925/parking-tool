"""
提案データJSONからPowerPointを生成するスクリプト
使い方: python generate_pptx.py parking_proposal_data.json
"""
import json, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

GREEN = RGBColor(0x1D, 0x9E, 0x75)
DARK = RGBColor(0x1a, 0x1a, 0x1a)
GRAY = RGBColor(0xf8, 0xf8, 0xf8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def add_slide(prs, layout_idx=6):
    return prs.slides.add_slide(prs.slide_layouts[layout_idx])

def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, text, x, y, w, h, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox

def add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def generate(data_path):
    with open(data_path) as f:
        data = json.load(f)

    a = data.get('analysis', {})
    m = data.get('market', {})
    fin = data.get('finance', {})
    prop = data.get('proposal', {})
    address = data.get('address', '東京都内')
    slides_data = prop.get('slides', [])

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    for s in slides_data:
        slide = add_slide(prs)
        set_bg(slide, WHITE)
        add_rect(slide, 0, 0, 13.33, 0.8, GREEN)
        add_text_box(slide, f"SLIDE {s['num']}  {s['title']}", 0.3, 0.15, 12, 0.6, size=20, bold=True, color=WHITE)
        add_rect(slide, 0, 0.8, 0.05, 6.7, GREEN)
        body = s.get('body', '')
        add_text_box(slide, body, 0.5, 1.1, 12.5, 5.8, size=14, color=DARK)
        if s['num'] == 6 and fin:
            add_text_box(slide, f"月間純利益: {fin.get('profit_monthly',0):,}円  /  年間: {fin.get('profit_yearly',0):,}円\n表面利回り: {fin.get('yield_rate',0)}%  /  回収: {fin.get('payback_months',0)}ヶ月", 0.5, 5.5, 12, 1.5, size=16, bold=True, color=GREEN)

    out = data_path.replace('.json', '.pptx')
    prs.save(out)
    print(f'✅ PowerPoint saved: {out}')

if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else 'parking_proposal_data.json'
    generate(path)
